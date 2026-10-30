"""Self-check for sheet reading: run `python test_import.py`."""

import os
import tempfile

from openpyxl import Workbook

from main import filterable_columns, import_excel_book

with tempfile.TemporaryDirectory() as tmp:
    path = os.path.join(tmp, "book.xlsx")
    wb = Workbook()

    # Clean sheet: repeated headers after trimming and a blank header
    ws = wb.active
    ws.title = "Clients"
    ws.append(["Client ID", "Client Name", "Email", "Email ", None])
    ws.append(["C1", "Acme", "a@acme.example", 1, "x"])
    ws.append([None, None, None, None, None])  # blank spacer row
    ws.append(["C2", "Beta", "b@beta.example", 2, "y"])

    # Title row above the real header: the "Unnamed: 0" case
    ws2 = wb.create_sheet("Titled")
    ws2.append(["ES Network Client Database"])
    ws2.append([])
    ws2.append(["Client ID", "Company", "Monthly Fee (PHP)", "Last Contact"])
    ws2.append(["CL-001", "Acme", 7500, "2026-09-24"])
    ws2.append(["CL-002", "Beta", 135500, "2026-09-28"])

    # Merged header cell: label must be spread across its columns
    ws3 = wb.create_sheet("Merged")
    ws3.append(["Contact", None, "Email"])
    ws3.merge_cells("A1:B1")
    ws3.append(["Andrei", "Aquino", "andrei@acme.example"])

    # No header at all: values must be kept, columns get numbered
    ws4 = wb.create_sheet("DataOnly")
    ws4.append(["a@acme.example", 7500, "2026-09-24"])
    ws4.append(["b@beta.example", 135500, "2026-09-28"])

    # Leading blank header cell: the literal "Unnamed: 0" case
    ws5 = wb.create_sheet("LeadingBlank")
    ws5.append([None, "Client", "Email"])
    ws5.append(["CL-001", "Acme", "a@acme.example"])

    wb.create_sheet("Empty")
    wb.save(path)

    book = import_excel_book(path)

by_sheet = {s["sheet"]: s for s in book}
assert list(by_sheet) == ["Clients", "Titled", "Merged", "DataOnly", "LeadingBlank"], list(by_sheet)

clients = by_sheet["Clients"]
assert clients["columns"] == ["Client ID", "Client Name", "Email", "Email (2)", "Column 5"], clients["columns"]
assert [r["row_index"] for r in clients["rows"]] == [1, 3], clients["rows"]
assert clients["rows"][1]["data"]["Client Name"] == "Beta"

titled = by_sheet["Titled"]
assert titled["header_row"] == 3, titled["header_row"]
assert titled["columns"] == ["Client ID", "Company", "Monthly Fee (PHP)", "Last Contact"], titled["columns"]
assert len(titled["rows"]) == 2, titled["rows"]  # title and blank row are not records

merged = by_sheet["Merged"]
assert merged["columns"] == ["Contact", "Contact (2)", "Email"], merged["columns"]

data_only = by_sheet["DataOnly"]
assert data_only["header_row"] is None
assert data_only["columns"] == ["Column 1", "Column 2", "Column 3"], data_only["columns"]
assert len(data_only["rows"]) == 2

assert by_sheet["Clients"]["rows"][0]["data"]["Email"] == "a@acme.example"

leading = by_sheet["LeadingBlank"]
assert leading["columns"] == ["Column 1", "Client", "Email"], leading["columns"]
assert not any(c.lower().startswith("unnamed") for s in book for c in s["columns"])

print("ok: all sheets, title rows skipped, merged headers filled, no Unnamed columns")

# Filter picker: only real, useful category columns
rows = [
    {
        "Status": "Active" if i % 2 else "Inactive",
        "Owner": "Rico" if i % 3 else "Dianne",
        "Notes": f"note {i}",          # too many distinct values
        "MostlyEmpty": "x" if i == 0 else "",
        "Single": "same",
        "Unnamed: 0": i,
        "Column 2": "junk",
    }
    for i in range(20)
]
picked = filterable_columns(rows)
assert set(picked) == {"Status", "Owner"}, picked
assert picked["Status"] == ["Active", "Inactive"]

assert filterable_columns([]) == {}
print("ok: filter picker skips numbered, single-value, mostly-empty and high-cardinality columns")

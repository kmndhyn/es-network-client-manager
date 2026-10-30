from fastapi import FastAPI, UploadFile, File, HTTPException, Query, Request
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
import pandas as pd
import openpyxl
import os
import json
import re
import shutil
from datetime import datetime
from typing import Optional, List, Dict, Any

from sqlalchemy import or_

from database import init_db, get_session, ImportMeta, ClientRow
from paths import STATIC_DIR, UPLOAD_DIR, DATA_DIR

app = FastAPI(title="ES Network Client Manager")

# Mount static files
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

init_db()

os.makedirs(UPLOAD_DIR, exist_ok=True)


# ─── Helpers ────────────────────────────────────────────────────────────────

NULLISH = {
    "", "nan", "none", "null", "nil", "n/a", "na", "-", "--", "---",
    "unknown", "unnamed", "undefined", "missing", "#n/a", "#na", "#null",
    "#ref!", "#value!", "#div/0!", "#name?", "no data", "not available",
    "—", "–", "•", "·",
}


def _clean_value(val):
    """Normalize pandas/numpy values and common null-placeholder strings to empty string."""
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except (TypeError, ValueError):
        pass  # list-like values
    if hasattr(val, "item") and not isinstance(val, (str, bytes)):
        val = val.item()
    if isinstance(val, (int, float, bool)):
        return val
    if isinstance(val, (pd.Timestamp, datetime)):
        return val.isoformat()
    if isinstance(val, str):
        stripped = val.strip()
        return "" if stripped.lower() in NULLISH else stripped
    return str(val)


def unique_columns(raw_columns) -> List[str]:
    """Every column survives: blanks get a name, repeats get a suffix."""
    names, used = [], set()
    for i, c in enumerate(raw_columns, start=1):
        name = str(c).strip()
        if not name or name.lower().startswith("unnamed:"):
            name = f"Column {i}"
        candidate, n = name, 1
        while candidate in used:
            n += 1
            candidate = f"{name} ({n})"
        used.add(candidate)
        names.append(candidate)
    return names


def _is_blank(val) -> bool:
    if val is None:
        return True
    try:
        if pd.isna(val):
            return True
    except (TypeError, ValueError):
        return False
    return str(val).strip() == ""


def _looks_like_data(val) -> bool:
    """Emails, links, numbers, dates and long text are data, not header labels."""
    if isinstance(val, bool):
        return False
    if isinstance(val, (int, float, datetime, pd.Timestamp)):
        return True
    text = str(val).strip()
    if "@" in text or text.lower().startswith(("http://", "https://", "www.")):
        return True
    if len(text) > 60:
        return True
    return bool(re.fullmatch(r"[\d\s,./:+-]+", text)) and any(ch.isdigit() for ch in text)


def _choose_header(df: pd.DataFrame, limit: int = 20) -> int:
    """Return the row index holding the real column names, or -1 when there is none.

    Scans the top of the sheet so a title/notes row above the table (the usual
    source of "Unnamed: 0") is skipped instead of becoming the header.
    """
    best, best_score = -1, 0
    for i in range(min(limit, len(df))):
        values = [v for v in df.iloc[i] if not _is_blank(v)]
        if len(values) < 2 or any(_looks_like_data(v) for v in values):
            continue
        if i + 1 < len(df) and all(_is_blank(v) for v in df.iloc[i + 1]):
            continue
        if len(values) > best_score:
            best, best_score = i, len(values)
    return best


def _header_names(df: pd.DataFrame, header_pos: int, ws, width: int) -> List[str]:
    """Header labels for one row, with merged ranges spread across their span."""
    if ws is not None:
        excel_row = header_pos + 1
        anchors = {}
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= excel_row <= rng.max_row:
                anchor = _clean_value(ws.cell(row=rng.min_row, column=rng.min_col).value)
                if anchor != "":
                    for col in range(rng.min_col, rng.max_col + 1):
                        anchors[col] = anchor
        raw = [
            anchors.get(i + 1, _clean_value(ws.cell(row=excel_row, column=i + 1).value))
            for i in range(width)
        ]
    else:
        raw = [_clean_value(v) for v in df.iloc[header_pos]]
        raw = (raw + [""] * width)[:width]
    return unique_columns(raw)


def import_excel_book(filepath: str) -> List[Dict[str, Any]]:
    """Read every sheet in a workbook, using each sheet's real header row."""
    book = pd.read_excel(filepath, sheet_name=None, header=None)
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
    except Exception:
        wb = None  # legacy .xls: header detection still works, merged labels are lost

    sheets = []
    for sheet_name, df in book.items():
        width = df.shape[1]
        if width == 0:
            continue  # genuinely empty sheet

        header_pos = _choose_header(df)
        ws = wb[sheet_name] if wb is not None else None
        if header_pos >= 0:
            columns = _header_names(df, header_pos, ws, width)
            body = df.iloc[header_pos + 1:]
        else:
            columns = [f"Column {i}" for i in range(1, width + 1)]
            body = df

        body = body.dropna(how="all")  # blank spacer rows carry no data
        rows = [
            {
                "row_index": int(idx),
                "data": {name: _clean_value(v) for name, v in zip(columns, values)},
            }
            for idx, values in zip(body.index, body.itertuples(index=False, name=None))
        ]
        sheets.append({
            "sheet": str(sheet_name),
            "columns": columns,
            "rows": rows,
            "header_row": header_pos + 1 if header_pos >= 0 else None,
        })

    if wb is not None:
        wb.close()
    return sheets


FILTER_LIMIT = 15
NUMBERED_COLUMN = re.compile(r"^(unnamed:\s*\d+|column\s+\d+)$", re.I)


def filterable_columns(all_data: List[Dict[str, Any]]) -> Dict[str, List[str]]:
    """Columns worth a dropdown: a handful of values, mostly filled, not a numbered placeholder."""
    filterable = {}
    total = len(all_data)
    if not total:
        return filterable

    for key in sorted({k for d in all_data for k in d}):
        if NUMBERED_COLUMN.match(str(key)):
            continue
        values = {str(d.get(key, "")).strip() for d in all_data}
        values.discard("")
        filled = sum(1 for d in all_data if str(d.get(key, "")).strip())
        if 2 <= len(values) <= FILTER_LIMIT and filled * 2 >= total:
            filterable[key] = sorted(values)
    return filterable


def rows_to_df(rows: List[ClientRow], columns: List[str]) -> pd.DataFrame:
    """Rebuild a sheet from stored rows: original order, manual rows last."""
    columns = list(columns)
    rows = sorted(rows, key=lambda r: (r.row_index if r.row_index is not None and r.row_index >= 0 else 10**9, r.id))
    parsed = [json.loads(r.data) if r.data else {} for r in rows]
    for d in parsed:
        for k in d:
            if k not in columns:
                columns.append(k)
    return pd.DataFrame({c: [d.get(c, "") for d in parsed] for c in columns})


def row_to_dict(row: ClientRow) -> Dict[str, Any]:
    return {
        "id": row.id,
        "source_file": row.source_file,
        "row_index": row.row_index,
        "data": json.loads(row.data) if row.data else {},
        "created_at": row.created_at.isoformat() if row.created_at else None,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
    }


# ─── Endpoints ──────────────────────────────────────────────────────────────

@app.get("/", response_class=HTMLResponse)
async def root():
    return FileResponse(str(STATIC_DIR / "index.html"))


@app.post("/api/import")
async def import_excel(files: List[UploadFile] = File(...)):
    if not files:
        raise HTTPException(status_code=400, detail="Upload at least 1 Excel file.")

    session = get_session()
    imported_total = 0
    sheet_total = 0
    no_header = []

    try:
        for upload in files:
            # Untrusted client filename: keep it inside uploads/
            filename = os.path.basename(upload.filename or "uploaded_file")
            dest = os.path.join(UPLOAD_DIR, filename)
            with open(dest, "wb") as f:
                shutil.copyfileobj(upload.file, f)

            try:
                sheets = import_excel_book(dest)
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Could not read {filename}: {e}")

            if not sheets:
                raise HTTPException(status_code=400, detail=f"{filename} has no data in any sheet.")

            # Re-importing a workbook replaces every sheet it had, including
            # rows from older versions that stored only the bare filename.
            prefix = filename + "::"
            for model in (ClientRow, ImportMeta):
                session.query(model).filter(
                    or_(
                        model.source_file == filename,
                        model.source_file.startswith(prefix, autoescape=True),
                    )
                ).delete(synchronize_session=False)

            for sheet in sheets:
                key = f"{filename}::{sheet['sheet']}"
                session.add(ImportMeta(source_file=key, columns=json.dumps(sheet["columns"])))
                if sheet["header_row"] is None:
                    no_header.append(f"{filename} › {sheet['sheet']}")
                for row in sheet["rows"]:
                    session.add(ClientRow(
                        source_file=key,
                        row_index=row["row_index"],
                        data=json.dumps(row["data"]),
                    ))
                    imported_total += 1
                sheet_total += 1

        session.commit()
    finally:
        session.close()

    message = f"Imported {imported_total} rows from {sheet_total} sheet(s) in {len(files)} file(s)."
    if no_header:
        message += " No header row found in: " + ", ".join(no_header) + " — columns are numbered."
    return {"imported": imported_total, "message": message}


@app.get("/api/clients")
async def list_clients(
    request: Request,
    search: Optional[str] = Query(None),
    source_file: Optional[str] = Query(None),
):
    session = get_session()
    q = session.query(ClientRow)
    if source_file:
        q = q.filter(ClientRow.source_file == source_file)
    rows = q.order_by(ClientRow.row_index).all()

    # Every other query param is a column filter from the UI's filter bar.
    filters = {k: v for k, v in request.query_params.items() if k not in ("search", "source_file")}

    results = []
    for row in rows:
        item = row_to_dict(row)
        if search:
            # Search across all data values
            haystack = " ".join(str(v) for v in item["data"].values()).lower()
            if search.lower() not in haystack:
                continue
        if any(str(item["data"].get(k, "")) != v for k, v in filters.items()):
            continue
        results.append(item)

    session.close()
    return results


@app.get("/api/columns")
async def get_columns(source_file: Optional[str] = Query(None)):
    """Return column definitions for the UI."""
    session = get_session()
    metas = session.query(ImportMeta).order_by(ImportMeta.id).all()
    sources = [m.source_file for m in metas]

    if source_file:
        meta = next((m for m in metas if m.source_file == source_file), None)
        if meta:
            columns = json.loads(meta.columns)
            session.close()
            return {"columns": columns, "source_file": source_file, "sources": sources}

    # Return all unique columns across all imports
    all_cols = []
    seen = set()
    for m in metas:
        for c in json.loads(m.columns):
            if c not in seen:
                all_cols.append(c)
                seen.add(c)
    session.close()
    return {"columns": all_cols, "source_file": source_file, "sources": sources}


@app.get("/api/filters")
async def get_filters(source_file: Optional[str] = Query(None)):
    """Return the columns worth showing as filter dropdowns."""
    session = get_session()
    q = session.query(ClientRow)
    if source_file:
        q = q.filter(ClientRow.source_file == source_file)
    all_data = [json.loads(r.data) or {} for r in q.all()]
    session.close()
    return {"filterable": filterable_columns(all_data)}


@app.post("/api/clients")
async def add_client(row: Dict[str, Any]):
    session = get_session()
    # Determine source_file from existing metas, or first available
    meta = session.query(ImportMeta).first()
    source_file = row.get("source_file") or (meta.source_file if meta else "manual")

    data = {k: v for k, v in row.items() if k not in ("source_file",)}
    client_row = ClientRow(
        source_file=source_file,
        row_index=-1,  # manually added
        data=json.dumps(data),
    )
    session.add(client_row)
    session.flush()
    client_id = client_row.id
    session.commit()
    session.close()
    return {"id": client_id, "message": "Client added."}


@app.put("/api/clients/{client_id}")
async def update_client(client_id: int, row: Dict[str, Any]):
    session = get_session()
    c = session.query(ClientRow).filter(ClientRow.id == client_id).first()
    if not c:
        session.close()
        raise HTTPException(status_code=404, detail="Client not found.")

    current_data = json.loads(c.data) if c.data else {}
    for key, value in row.items():
        if key not in ("source_file", "id"):
            current_data[key] = value
    c.data = json.dumps(current_data)
    session.commit()
    session.close()
    return {"message": "Client updated."}


@app.delete("/api/clients/{client_id}")
async def delete_client(client_id: int):
    session = get_session()
    c = session.query(ClientRow).filter(ClientRow.id == client_id).first()
    if not c:
        session.close()
        raise HTTPException(status_code=404, detail="Client not found.")
    session.delete(c)
    session.commit()
    session.close()
    return {"message": "Client deleted."}


@app.post("/api/import/reset")
async def reset_import():
    """Clear database and uploaded files. Full reset for a fresh import."""
    session = get_session()
    session.query(ClientRow).delete()
    session.query(ImportMeta).delete()
    session.commit()
    session.close()
    for fname in os.listdir(UPLOAD_DIR):
        os.remove(os.path.join(UPLOAD_DIR, fname))
    return {"message": "Database cleared. Ready for a fresh import."}


@app.post("/api/save-to-excel")
async def save_to_excel():
    """Write current database back to the stored Excel files, one sheet per dataset."""
    session = get_session()
    all_rows = session.query(ClientRow).order_by(ClientRow.id).all()
    metas = {m.source_file: json.loads(m.columns or "[]") for m in session.query(ImportMeta).all()}
    session.close()

    if not all_rows:
        raise HTTPException(status_code=400, detail="No data to save.")

    # Group rows by workbook, then by sheet
    by_file: Dict[str, Dict[Optional[str], List[ClientRow]]] = {}
    for row in all_rows:
        fname, _, sheet = (row.source_file or "").partition("::")
        by_file.setdefault(fname, {}).setdefault(sheet or None, []).append(row)

    saved = []
    failed = []
    for fname, sheet_rows in by_file.items():
        fpath = os.path.join(UPLOAD_DIR, fname)
        if not os.path.exists(fpath):
            continue

        try:
            # Legacy rows (imported before per-sheet storage) have no sheet name
            if None in sheet_rows:
                first_sheet = pd.ExcelFile(fpath).sheet_names[0]
                sheet_rows[first_sheet] = sheet_rows.pop(None)
                metas.setdefault(f"{fname}::{first_sheet}", metas.get(fname, []))

            with pd.ExcelWriter(fpath, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                for sheet, rows in sheet_rows.items():
                    df = rows_to_df(rows, metas.get(f"{fname}::{sheet}", []))
                    df.to_excel(writer, sheet_name=sheet, index=False)
            saved.append(fname)
        except Exception as e:
            failed.append(f"{fname} ({e})")

    if saved:
        message = f"Saved {len(saved)} file(s): {', '.join(saved)}"
    else:
        message = "No files saved. Make sure files are still in uploads."
    if failed:
        message += f" · Could not save: {'; '.join(failed)}"
    return {"message": message}


@app.get("/api/export")
async def export_excel():
    """Download the whole database as one sheet per imported sheet."""
    session = get_session()
    rows = session.query(ClientRow).order_by(ClientRow.id).all()
    metas = {m.source_file: json.loads(m.columns or "[]") for m in session.query(ImportMeta).order_by(ImportMeta.id).all()}
    session.close()

    if not rows:
        raise HTTPException(status_code=400, detail="No data to export.")

    by_source: Dict[str, List[ClientRow]] = {}
    for row in rows:
        by_source.setdefault(row.source_file or "manual", []).append(row)

    # Keep workbook/sheet order, then append anything without metadata
    order = [k for k in metas if k in by_source] + [k for k in by_source if k not in metas]
    used_names, sheets = set(), []
    for key in order:
        fname, _, sheet = key.partition("::")
        name = re.sub(r"[\[\]:*?/\\]", "", sheet or os.path.splitext(fname)[0])[:31] or "Data"
        base, n = name, 1
        while name.lower() in used_names:
            n += 1
            suffix = f" ({n})"
            name = base[:31 - len(suffix)] + suffix
        used_names.add(name.lower())
        sheets.append((name, rows_to_df(by_source[key], metas.get(key, []))))

    export_path = DATA_DIR / "export.xlsx"
    with pd.ExcelWriter(export_path, engine="openpyxl") as writer:
        for name, df in sheets:
            df.to_excel(writer, sheet_name=name, index=False)
    return FileResponse(str(export_path), filename="es_network_export.xlsx")
